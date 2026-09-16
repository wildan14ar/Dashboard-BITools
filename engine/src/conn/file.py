"""Konektor tabular read-only: file upload (CSV/XLSX), URL publik, Google Sheets.

Pola non-SQL seperti mongodb/api: field `sql` pada dataset dipakai sebagai
perintah JSON, mis. `{"sheet": "Sheet1", "limit": 1000}` (kosong = default).
Hasil selalu tampil mentah (list[dict]) — tanpa query SQL.

Arus file besar dibatasi berlapis (bukan dialirkan sampai browser):
- HTTP fetch di-cap FETCH_MAX_BYTES dan dibaca per chunk ke spool file
  (spill ke disk otomatis), bukan `resp.read()` sekaligus.
- Parse berhenti lebih awal (early-stop) begitu cap baris tercapai:
  cap = min(limit perintah, MAX_ROWS_HARD). Satu baris probe ekstra
  menentukan flag `last_truncated` yang dibaca executor.
- XLSX dibuka read_only + iter_rows (tidak membangun DOM penuh).
"""

from __future__ import annotations

import csv
import io
import ipaddress
import json
import logging
import re
import socket
import tempfile
import time
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path
from typing import Any

from src.config import (
    DATA_DIR,
    FETCH_MAX_BYTES,
    FILE_READ_CHUNK_SIZE,
    FILE_SCHEMA_SAMPLE_ROWS,
    MAX_ROWS_HARD,
    SHEETS_PAGE_ROWS,
)
from src.conn import Engine as BaseEngine
from src.conn import register

logger = logging.getLogger(__name__)

ALLOWED_EXTS = {".csv", ".xlsx"}
CSV_SNIFF_BYTES = 32 * 1024
SHEETS_ID_RE = re.compile(r"/d/([A-Za-z0-9-_]+)")
SHEETS_API_BASE = "https://sheets.googleapis.com/v4/spreadsheets"
SHEETS_SCOPE = "https://www.googleapis.com/auth/spreadsheets.readonly"


def _norm_cell(v: Any) -> str | int | float | bool | None:
    if v is None:
        return None
    if isinstance(v, (str, int, float, bool)):
        return v
    return str(v)


def _blank_name(idx: int) -> str:
    return f"column_{idx + 1}"


def _xlsx_cell(c: Any) -> Any:
    if c is None:
        return ""
    if isinstance(c, bool):
        return c
    if isinstance(c, (int, float)):
        return c
    if isinstance(c, str):
        return c.strip()
    return str(c)


def _xlsx_blank(c: Any) -> bool:
    return c is None or (isinstance(c, str) and not c.strip()) or c == ""


def _infer_type(samples: list[str]) -> str:
    vals = [s for s in samples if s not in ("", None)]
    if not vals:
        return "TEXT"
    try:
        for v in vals:
            int(str(v))
        return "INTEGER"
    except (ValueError, TypeError):
        pass
    try:
        for v in vals:
            float(str(v))
        return "REAL"
    except (ValueError, TypeError):
        pass
    return "TEXT"


class _Deadline:
    def __init__(self, timeout_sec: int | None):
        self._at = time.monotonic() + timeout_sec if timeout_sec else None

    def check(self, what: str = "file operation"):
        if self._at is not None and time.monotonic() > self._at:
            raise TimeoutError(f"{what} exceeded timeout")


# ---------------------------------------------------------------------------
# SSRF guard untuk fetch URL
# ---------------------------------------------------------------------------


def _assert_public_http(url: str) -> urllib.parse.ParseResult:
    parts = urllib.parse.urlparse(url)
    if parts.scheme not in ("http", "https"):
        raise ValueError("URL harus http(s)")
    if parts.username or parts.password:
        raise ValueError("URL tidak boleh memuat kredensial")
    host = parts.hostname or ""
    try:
        infos = socket.getaddrinfo(host, None)
    except OSError as e:
        raise ValueError(f"Host tidak dapat di-resolve: {host}") from e
    for info in infos:
        ip = ipaddress.ip_address(info[4][0])
        if (
            ip.is_private
            or ip.is_loopback
            or ip.is_link_local
            or ip.is_multicast
            or ip.is_reserved
            or ip.is_unspecified
        ):
            raise ValueError(f"Host {host} menunjuk ke alamat privat/dilarang")
    return parts


class _GuardedRedirectHandler(urllib.request.HTTPRedirectHandler):
    """Ikuti redirect hanya ke target http(s) publik, maksimal 5 hop."""

    max_redirections = 5

    def redirect_request(self, req, fp, code, msg, headers, newurl):
        _assert_public_http(urllib.parse.urljoin(req.full_url, newurl))
        return super().redirect_request(req, fp, code, msg, headers, newurl)


def _fetch_to_spool(
    url: str, timeout: float, max_bytes: int, deadline: _Deadline
) -> tuple[tempfile.SpooledTemporaryFile, str | None]:
    """GET streaming ke spool file (spill ke disk >10MB). Kembalikan (file, content_type)."""
    _assert_public_http(url)
    opener = urllib.request.build_opener(_GuardedRedirectHandler())
    req = urllib.request.Request(url, headers={"User-Agent": "BITools-FileEngine/1.0"})
    try:
        resp = opener.open(req, timeout=timeout)
    except urllib.error.HTTPError as e:
        raise RuntimeError(f"Fetch {url} gagal: HTTP {e.code}") from e
    except urllib.error.URLError as e:
        raise ConnectionError(f"Tidak dapat menjangkau {url}: {e.reason}") from e
    try:
        length = resp.headers.get("Content-Length")
        if length is not None and length.isdigit() and int(length) > max_bytes:
            raise ValueError(f"Ukuran file ({length} bytes) melebihi batas {max_bytes} bytes")
        spool: tempfile.SpooledTemporaryFile = tempfile.SpooledTemporaryFile(  # noqa: SIM115 - dikembalikan ke pemanggil
            max_size=10 * 1024 * 1024, mode="w+b"
        )
        total = 0
        while True:
            deadline.check("fetch URL")
            chunk = resp.read(FILE_READ_CHUNK_SIZE)
            if not chunk:
                break
            total += len(chunk)
            if total > max_bytes:
                raise ValueError(f"Ukuran file melebihi batas {max_bytes} bytes")
            spool.write(chunk)
        spool.seek(0)
        return spool, resp.headers.get_content_type()
    finally:
        resp.close()


# ---------------------------------------------------------------------------
# Parser streaming
# ---------------------------------------------------------------------------


def _sniff_dialect(sample: bytes):
    try:
        dialect = csv.Sniffer().sniff(sample.decode("utf-8-sig", errors="replace"))
        delim = getattr(dialect, "delimiter", ",")
        if not isinstance(delim, str) or len(delim) != 1 or not delim.isprintable():
            return csv.excel
        return dialect
    except Exception:
        return csv.excel


def _iter_csv_rows(
    raw: io.BufferedReader | io.RawIOBase, deadline: _Deadline
) -> tuple[list[str], Any]:
    """Kembalikan (header, generator baris-dict). Streaming, tanpa buffering penuh."""
    sample = raw.read(CSV_SNIFF_BYTES)
    dialect = _sniff_dialect(sample)
    raw.seek(0)
    text = io.TextIOWrapper(raw, encoding="utf-8-sig", errors="replace", newline="")
    reader = csv.reader(text, dialect)
    try:
        first = next(reader)
    except StopIteration:
        return [], iter(())
    header = [c.strip() or _blank_name(i) for i, c in enumerate(first)]

    def gen():
        for row in reader:
            deadline.check("parse CSV")
            if not any((c or "").strip() for c in row):
                continue
            yield {h: _norm_cell(row[i] if i < len(row) else "") for i, h in enumerate(header)}

    return header, gen()


def _open_xlsx_sheet(
    raw: io.BufferedReader | io.RawIOBase | str | Path, sheet: str | None
):
    from openpyxl import load_workbook

    wb = load_workbook(raw, read_only=True, data_only=True)
    names = wb.sheetnames
    if not names:
        raise ValueError("Workbook tidak memiliki sheet")
    target = sheet or names[0]
    if target not in names:
        raise ValueError(f"Sheet {target!r} tidak ditemukan. Tersedia: {', '.join(names)}")
    return wb, wb[target], names


def _iter_xlsx_rows(ws, deadline: _Deadline) -> tuple[list[str], Any]:
    rows = ws.iter_rows(values_only=True)
    try:
        first = next(rows)
    except StopIteration:
        return [], iter(())
    header = []
    for i, c in enumerate(first or []):
        s = "" if c is None else str(c).strip()
        header.append(s or _blank_name(i))

    def gen():
        for row in rows:
            deadline.check("parse XLSX")
            cells = [_xlsx_cell(c) for c in (row or [])]
            if all(_xlsx_blank(c) for c in cells):
                continue
            yield {h: _norm_cell(cells[i] if i < len(cells) else "") for i, h in enumerate(header)}

    return header, gen()


def _detect_format(name: str, content_type: str | None) -> str:
    low = name.lower()
    if low.endswith(".xlsx"):
        return "xlsx"
    if low.endswith(".csv"):
        return "csv"
    if content_type in ("text/csv", "application/csv"):
        return "csv"
    if content_type in (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ):
        return "xlsx"
    raise ValueError(f"Format file {name!r} tidak didukung (hanya .csv/.xlsx)")


# ---------------------------------------------------------------------------
# Google Sheets
# ---------------------------------------------------------------------------


def _sheets_id(value: str) -> str:
    value = (value or "").strip()
    m = SHEETS_ID_RE.search(value)
    if m:
        return m.group(1)
    if re.fullmatch(r"[A-Za-z0-9-_]+", value):
        return value
    raise ValueError("spreadsheet_id tidak valid (ID atau share-link Sheets)")


def _sheets_export_url(spreadsheet_id: str, gid: str | None) -> str:
    url = f"https://docs.google.com/spreadsheets/d/{spreadsheet_id}/export?format=csv"
    if gid:
        if not gid.isdigit():
            raise ValueError("gid harus numerik")
        url += f"&gid={gid}"
    return url


def _sheets_token(service_account_json: str) -> str:
    from google.auth.transport.requests import Request
    from google.oauth2 import service_account

    try:
        info = json.loads(service_account_json)
    except json.JSONDecodeError as e:
        raise ValueError(f"service_account_json bukan JSON valid: {e}") from e
    creds = service_account.Credentials.from_service_account_info(info, scopes=[SHEETS_SCOPE])
    creds.refresh(Request())
    if not creds.token:
        raise RuntimeError("Gagal memperoleh access token service account")
    return creds.token


def _sheets_get_json(url: str, token: str | None, api_key: str | None, timeout: float) -> Any:
    if api_key and "key=" not in url:
        url += ("&" if "?" in url else "?") + "key=" + urllib.parse.quote(api_key)
    headers = {"Accept": "application/json"}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    _assert_public_http(url)
    req = urllib.request.Request(url, headers=headers)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            return json.loads(resp.read().decode("utf-8") or "null")
    except urllib.error.HTTPError as e:
        raise RuntimeError(f"Sheets API error {e.code}: {e.reason}") from e
    except urllib.error.URLError as e:
        raise ConnectionError(f"Tidak dapat menjangkau Sheets API: {e.reason}") from e


def _sheets_props(spreadsheet_id: str, token: str | None, api_key: str | None, timeout: float) -> list[tuple[int, str]]:
    meta = _sheets_get_json(
        f"{SHEETS_API_BASE}/{spreadsheet_id}?fields=sheets.properties",
        token,
        api_key,
        timeout,
    )
    out = []
    for s in meta.get("sheets") or []:
        props = s.get("properties") or {}
        if props.get("title") is not None and props.get("sheetId") is not None:
            out.append((int(props["sheetId"]), str(props["title"])))
    return out


def _sheets_titles(spreadsheet_id: str, token: str | None, api_key: str | None, timeout: float) -> list[str]:
    return [title for _, title in _sheets_props(spreadsheet_id, token, api_key, timeout)]


def _iter_sheets_values(
    spreadsheet_id: str,
    title: str,
    token: str | None,
    api_key: str | None,
    timeout: float,
    deadline: _Deadline,
) -> tuple[list[str], Any]:
    """Paging rentang per SHEETS_PAGE_ROWS baris; kembalikan (header, generator)."""
    header: list[str] = []
    header_set = False
    buf: list[dict] = []
    start = 1
    exhausted = False

    def gen():
        nonlocal start, exhausted, header_set
        while not exhausted:
            deadline.check("fetch Google Sheets")
            end = start + SHEETS_PAGE_ROWS - 1
            payload = _sheets_get_json(
                f"{SHEETS_API_BASE}/{spreadsheet_id}/values/"
                f"{urllib.parse.quote(title, safe='')}!{start}:{end}"
                "?majorDimension=ROWS",
                token,
                api_key,
                timeout,
            )
            values = payload.get("values") or []
            if not values:
                exhausted = True
                return
            for row in values:
                cells = [str(c).strip() for c in row]
                if not header_set:
                    header.extend([c or _blank_name(i) for i, c in enumerate(cells)])
                    header_set = True
                    continue
                if not any(cells):
                    continue
                buf.append(
                    {h: _norm_cell(cells[i] if i < len(cells) else "") for i, h in enumerate(header)}
                )
                if buf:
                    yield buf.pop(0)
            if len(values) < SHEETS_PAGE_ROWS:
                exhausted = True
            start = end + 1
        while buf:
            yield buf.pop(0)

    g = gen()
    try:
        first = next(g)
        h = list(header)

        def replay():
            yield first
            yield from g

        return h, replay()
    except StopIteration:
        return list(header), iter(())


# ---------------------------------------------------------------------------
# Engine
# ---------------------------------------------------------------------------


def _parse_command(sql: str) -> dict:
    sql = (sql or "").strip()
    if not sql:
        return {}
    try:
        cmd = json.loads(sql)
    except json.JSONDecodeError as e:
        raise ValueError(
            'Perintah file harus JSON, mis. {"sheet": "Sheet1", "limit": 1000}'
        ) from e
    if not isinstance(cmd, dict):
        raise TypeError("Perintah file harus objek JSON")
    if "limit" in cmd:
        try:
            limit = int(cmd["limit"])
        except (TypeError, ValueError):
            raise ValueError("limit harus bilangan bulat positif") from None
        if limit <= 0:
            raise ValueError("limit harus bilangan bulat positif")
        cmd["limit"] = limit
    return cmd


@register("file")
class FileEngine(BaseEngine):
    """Konektor tabular read-only: upload (CSV/XLSX), URL publik, Google Sheets."""

    def __init__(self, source_id: str, config: dict) -> None:
        self._source_id = source_id
        self._config = dict(config or {})
        self._kind = str(self._config.get("kind", "upload")).lower()
        if self._kind not in ("upload", "url", "sheets"):
            raise ValueError(f"kind file tidak dikenal: {self._kind!r}")
        self.last_truncated = False
        logger.info(f"Created file engine for [{source_id}] kind={self._kind}")

    # -- sumber bytes -----------------------------------------------------

    def _resolve_upload(self) -> tuple[Path, str]:
        rel = str(self._config.get("path", "")).strip()
        if not rel:
            raise ValueError("Source file (upload) membutuhkan 'path'")
        base = Path(DATA_DIR).resolve()
        target = (base / rel).resolve()
        if target != base and base not in target.parents:
            raise ValueError("Path upload di luar DATA_DIR")
        if not target.is_file():
            raise ValueError(f"File upload tidak ditemukan: {rel}")
        ext = target.suffix.lower()
        if ext not in ALLOWED_EXTS:
            raise ValueError(f"Ekstensi {ext!r} tidak didukung (hanya .csv/.xlsx)")
        return target, ext[1:]

    def _open_upload(self, deadline: _Deadline):
        target, fmt = self._resolve_upload()
        deadline.check("open upload")
        return open(target, "rb"), target.name, fmt

    def _open_url(self, deadline: _Deadline, timeout: float):
        url = str(self._config.get("file_url", "")).strip()
        if not url:
            raise ValueError("Source file (url) membutuhkan 'file_url'")
        spool, content_type = _fetch_to_spool(url, timeout, FETCH_MAX_BYTES, deadline)
        name = urllib.parse.urlparse(url).path.rsplit("/", 1)[-1] or "remote"
        fmt = str(self._config.get("format", "")).lower() or _detect_format(name, content_type)
        if fmt not in ("csv", "xlsx"):
            raise ValueError(f"Format {fmt!r} tidak didukung (hanya csv/xlsx)")
        return spool, name, fmt

    def _sheets_params(self) -> dict:
        cfg = self._config
        spreadsheet_id = _sheets_id(str(cfg.get("spreadsheet_id", "")))
        auth = str(cfg.get("auth", "none")).lower()
        if auth not in ("none", "api_key", "service_account"):
            raise ValueError(f"auth sheets tidak dikenal: {auth!r}")
        api_key = str(cfg.get("api_key", "") or "").strip() or None
        sa_json = str(cfg.get("service_account_json", "") or "").strip() or None
        if auth == "api_key" and not api_key:
            raise ValueError("auth=api_key membutuhkan 'api_key'")
        if auth == "service_account" and not sa_json:
            raise ValueError("auth=service_account membutuhkan 'service_account_json'")
        return {
            "spreadsheet_id": spreadsheet_id,
            "sheet": str(cfg.get("sheet", "") or "").strip() or None,
            "gid": str(cfg.get("gid", "") or "").strip() or None,
            "auth": auth,
            "api_key": api_key,
            "sa_json": sa_json,
        }

    def _sheets_token(self, p: dict) -> str | None:
        if p["auth"] == "service_account" and p["sa_json"]:
            return _sheets_token(p["sa_json"])
        return None

    # -- Engine API --------------------------------------------------------

    def _read_table(
        self, sheet: str | None, row_cap: int, timeout: float, deadline: _Deadline
    ) -> tuple[str, list[str], list[dict], bool]:
        """Baca satu tabel s/d row_cap baris + 1 probe. Kembalikan (nama, header, rows, truncated)."""
        kind = self._kind
        if kind == "sheets":
            return self._read_sheets(sheet, row_cap, timeout, deadline)
        if kind == "url":
            raw, name, fmt = self._open_url(deadline, timeout)
        else:
            raw, name, fmt = self._open_upload(deadline)
        try:
            if fmt == "csv":
                header, gen = _iter_csv_rows(raw, deadline)
                table = Path(name).stem
            else:
                wb, ws, _names = _open_xlsx_sheet(raw, sheet)
                try:
                    header, gen = _iter_xlsx_rows(ws, deadline)
                    table = ws.title
                finally:
                    wb.close()
            rows: list[dict] = []
            for row in gen:
                rows.append(row)
                if len(rows) > row_cap:
                    return table, header, rows[:row_cap], True
            return table, header, rows, False
        finally:
            try:
                raw.close()
            except Exception:
                logger.debug("Could not close raw stream", exc_info=True)

    def _read_sheets(
        self, sheet: str | None, row_cap: int, timeout: float, deadline: _Deadline
    ) -> tuple[str, list[str], list[dict], bool]:
        p = self._sheets_params()
        token = self._sheets_token(p)
        api_key = p["api_key"]
        if p["auth"] == "none":
            spool, _ = _fetch_to_spool(
                _sheets_export_url(p["spreadsheet_id"], p["gid"]),
                timeout,
                FETCH_MAX_BYTES,
                deadline,
            )
            try:
                header, gen = _iter_csv_rows(spool, deadline)
                table = p["sheet"] or "Sheet1"
            except Exception:
                spool.close()
                raise
            rows: list[dict] = []
            try:
                for row in gen:
                    rows.append(row)
                    if len(rows) > row_cap:
                        return table, header, rows[:row_cap], True
                return table, header, rows, False
            finally:
                spool.close()
        # v4 API (api_key / service_account): paging per rentang.
        title = p["sheet"]
        if not title:
            if p["gid"]:
                props = _sheets_props(p["spreadsheet_id"], token, api_key, timeout)
                match = next((t for gid, t in props if str(gid) == p["gid"]), None)
                if match is None:
                    raise ValueError(f"gid {p['gid']!r} tidak ditemukan di spreadsheet")
                title = match
            else:
                titles = _sheets_titles(p["spreadsheet_id"], token, api_key, timeout)
                title = titles[0] if titles else "Sheet1"
        header, gen = _iter_sheets_values(
            p["spreadsheet_id"], title, token, api_key, timeout, deadline
        )
        rows = []
        for row in gen:
            rows.append(row)
            if len(rows) > row_cap:
                return title, header, rows[:row_cap], True
        return title, header, rows, False

    def fetch_all(
        self, sql: str, params: dict | None = None, timeout_sec: int | None = None
    ) -> list[dict]:
        self.last_truncated = False
        cmd = _parse_command(sql)
        timeout = float(timeout_sec or 30)
        deadline = _Deadline(timeout_sec or 30)
        row_cap = min(int(cmd.get("limit", MAX_ROWS_HARD)), MAX_ROWS_HARD)
        table, _header, rows, truncated = self._read_table(
            cmd.get("sheet"), row_cap, timeout, deadline
        )
        self.last_truncated = truncated
        logger.info(f"File read [{self._source_id}] table={table} rows={len(rows)} truncated={truncated}")
        return rows

    def list_tables(self) -> list[dict]:
        deadline = _Deadline(30)
        kind = self._kind
        if kind == "sheets":
            return self._sheets_tables(deadline)
        if kind == "url":
            raw, name, fmt = self._open_url(deadline, 30.0)
        else:
            raw, name, fmt = self._open_upload(deadline)
        try:
            if fmt == "csv":
                header, gen = _iter_csv_rows(raw, deadline)
                samples: dict[str, list[str]] = {h: [] for h in header}
                for row in gen:
                    for h in header:
                        if len(samples[h]) < FILE_SCHEMA_SAMPLE_ROWS:
                            samples[h].append(str(row.get(h, "")))
                    if all(len(v) >= FILE_SCHEMA_SAMPLE_ROWS for v in samples.values()):
                        break
                cols = [
                    {"name": h, "type": _infer_type(samples[h]), "nullable": True}
                    for h in header
                ]
                return [{"name": Path(name).stem, "schema": "", "columns": cols}]
            wb, _ws, names = _open_xlsx_sheet(raw, None)
            try:
                out = []
                for sheet_name in names:
                    ws = wb[sheet_name]
                    header, gen = _iter_xlsx_rows(ws, deadline)
                    samples = {h: [] for h in header}
                    for row in gen:
                        for h in header:
                            if len(samples[h]) < FILE_SCHEMA_SAMPLE_ROWS:
                                samples[h].append(str(row.get(h, "")))
                        if all(len(v) >= FILE_SCHEMA_SAMPLE_ROWS for v in samples.values()):
                            break
                    cols = [
                        {"name": h, "type": _infer_type(samples[h]), "nullable": True}
                        for h in header
                    ]
                    out.append({"name": sheet_name, "schema": "", "columns": cols})
                return out
            finally:
                wb.close()
        finally:
            try:
                raw.close()
            except Exception:
                logger.debug("Could not close raw stream", exc_info=True)

    def _sheets_tables(self, deadline: _Deadline) -> list[dict]:
        p = self._sheets_params()
        token = self._sheets_token(p)
        if p["auth"] == "none":
            spool, _ = _fetch_to_spool(
                _sheets_export_url(p["spreadsheet_id"], p["gid"]), 30.0, FETCH_MAX_BYTES, deadline
            )
            try:
                header, gen = _iter_csv_rows(spool, deadline)
                samples = {h: [] for h in header}
                for row in gen:
                    for h in header:
                        if len(samples[h]) < FILE_SCHEMA_SAMPLE_ROWS:
                            samples[h].append(str(row.get(h, "")))
                    if all(len(v) >= FILE_SCHEMA_SAMPLE_ROWS for v in samples.values()):
                        break
                cols = [
                    {"name": h, "type": _infer_type(samples[h]), "nullable": True}
                    for h in header
                ]
                return [{"name": p["sheet"] or "Sheet1", "schema": "", "columns": cols}]
            finally:
                spool.close()
        titles = _sheets_titles(p["spreadsheet_id"], token, p["api_key"], 30.0)
        out = []
        for title in titles:
            header, gen = _iter_sheets_values(
                p["spreadsheet_id"], title, token, p["api_key"], 30.0, deadline
            )
            samples = {h: [] for h in header}
            for row in gen:
                for h in header:
                    if len(samples[h]) < FILE_SCHEMA_SAMPLE_ROWS:
                        samples[h].append(str(row.get(h, "")))
                if all(len(v) >= FILE_SCHEMA_SAMPLE_ROWS for v in samples.values()):
                    break
            cols = [
                {"name": h, "type": _infer_type(samples[h]), "nullable": True} for h in header
            ]
            out.append({"name": title, "schema": "", "columns": cols})
        return out

    def close(self) -> None:
        pass

    def ping(self) -> bool:
        tables = self.list_tables()
        if not tables:
            raise ValueError("Tidak ada tabel/sheet terbaca dari source file")
        return True
