import csv
import json

import pytest

import src.conn.file as file_mod
from src import factory
from src.conn import create as conn_create
from src.executor import execute_query, get_schema_info


@pytest.fixture()
def datadir(tmp_path, monkeypatch):
    uploads = tmp_path / "uploads"
    uploads.mkdir()
    monkeypatch.setattr(file_mod, "DATA_DIR", str(tmp_path))
    return tmp_path


def _write_csv(path, rows, header=("nama", "umur")):
    with open(path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


def _upload_config(name):
    return json.dumps({"kind": "upload", "path": f"uploads/{name}"})


def test_csv_upload_read(datadir):
    _write_csv(datadir / "uploads" / "data.csv", [[f"u{i}", i] for i in range(5)])
    engine = conn_create("file", "f_csv", _upload_config("data.csv"))
    rows = engine.fetch_all("", {})
    assert len(rows) == 5
    assert rows[0] == {"nama": "u0", "umur": "0"}
    assert engine.last_truncated is False


def test_csv_limit_truncated_probe(datadir):
    _write_csv(datadir / "uploads" / "big.csv", [[f"u{i}", i] for i in range(10)])
    engine = conn_create("file", "f_lim", _upload_config("big.csv"))
    rows = engine.fetch_all(json.dumps({"limit": 3}), {})
    assert len(rows) == 3
    assert engine.last_truncated is True


def test_csv_exact_limit_not_truncated(datadir):
    _write_csv(datadir / "uploads" / "exact.csv", [[f"u{i}", i] for i in range(3)])
    engine = conn_create("file", "f_exact", _upload_config("exact.csv"))
    rows = engine.fetch_all(json.dumps({"limit": 3}), {})
    assert len(rows) == 3
    assert engine.last_truncated is False


def test_xlsx_multisheet(datadir):
    from openpyxl import Workbook

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Satu"
    ws1.append(["a", "b"])
    ws1.append([1, 2])
    ws2 = wb.create_sheet("Dua")
    ws2.append(["x"])
    ws2.append(["y"])
    wb.save(datadir / "uploads" / "multi.xlsx")

    engine = conn_create("file", "f_xlsx", _upload_config("multi.xlsx"))
    default_rows = engine.fetch_all("", {})
    assert default_rows == [{"a": 1, "b": 2}]
    dua_rows = engine.fetch_all(json.dumps({"sheet": "Dua"}), {})
    assert dua_rows == [{"x": "y"}]
    tables = engine.list_tables()
    assert {t["name"] for t in tables} == {"Satu", "Dua"}
    assert [c["name"] for c in tables[0]["columns"]] == ["a", "b"]


def test_xlsx_missing_sheet(datadir):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.append(["a"])
    wb.save(datadir / "uploads" / "one.xlsx")
    engine = conn_create("file", "f_xlsx_miss", _upload_config("one.xlsx"))
    with pytest.raises(ValueError, match="tidak ditemukan"):
        engine.fetch_all(json.dumps({"sheet": "GaAda"}), {})


def test_invalid_command(datadir):
    _write_csv(datadir / "uploads" / "c.csv", [["a"]])
    engine = conn_create("file", "f_bad", _upload_config("c.csv"))
    with pytest.raises(ValueError, match="JSON"):
        engine.fetch_all("SELECT 1", {})


def test_upload_path_traversal_rejected():
    engine = conn_create("file", "f_trav", json.dumps({"kind": "upload", "path": "../secret.csv"}))
    with pytest.raises(ValueError, match="DATA_DIR"):
        engine.fetch_all("", {})


def test_upload_unsupported_extension(datadir):
    (datadir / "uploads" / "data.xls").write_text("x")
    engine = conn_create("file", "f_xls", _upload_config("data.xls"))
    with pytest.raises(ValueError, match="tidak didukung"):
        engine.fetch_all("", {})


def test_url_ssrf_guards():
    for bad in (
        "file:///etc/passwd",
        "http://127.0.0.1/data.csv",
        "http://localhost/data.csv",
        "http://169.254.169.254/latest",
        "https://10.0.0.1/x.csv",
    ):
        engine = conn_create("file", f"f_ssrf_{bad}", json.dumps({"kind": "url", "file_url": bad}))
        with pytest.raises((ValueError, ConnectionError, RuntimeError)):
            engine.fetch_all("", {})


def test_sheets_id_parsing():
    assert file_mod._sheets_id("1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms") == (
        "1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms"
    )
    assert (
        file_mod._sheets_id(
            "https://docs.google.com/spreadsheets/d/ABC123_xyz/edit#gid=0"
        )
        == "ABC123_xyz"
    )
    with pytest.raises(ValueError):
        file_mod._sheets_id("not a valid id!!")


def test_sheets_export_url():
    url = file_mod._sheets_export_url("ABC", "0")
    assert "export?format=csv" in url and "gid=0" in url
    url2 = file_mod._sheets_export_url("ABC", None)
    assert "gid" not in url2


def test_executor_truncated_flow(datadir):
    _write_csv(datadir / "uploads" / "flow.csv", [[f"u{i}", i] for i in range(10)])
    try:
        out = execute_query(
            source_id="flow_file",
            sql="",
            db_type="file",
            config_json=_upload_config("flow.csv"),
            max_rows=4,
            timeout_sec=10,
            use_cache=False,
        )
        assert out["row_count"] == 4
        assert out["truncated"] is True
        assert out["columns"] == ["nama", "umur"]
    finally:
        factory.dispose("flow_file")


def test_file_schema_info(datadir):
    _write_csv(datadir / "uploads" / "schema.csv", [["a", "1"], ["b", "2"]], header=("nama", "umur"))
    try:
        result = get_schema_info("schema_file", "file", _upload_config("schema.csv"))
        assert len(result["tables"]) == 1
        cols = {c["name"]: c["type"] for c in result["tables"][0]["columns"]}
        assert cols == {"nama": "TEXT", "umur": "INTEGER"}
    finally:
        factory.dispose("schema_file")


def test_file_ping(datadir):
    _write_csv(datadir / "uploads" / "ping.csv", [["a"]])
    assert file_mod.FileEngine("p", {"kind": "upload", "path": "uploads/ping.csv"}).ping() is True
